"""
Generic analysis-file builder
=============================
Any set of CSV / Excel files  ->  profile every table  ->  find the linking key  ->  decide the grain
->  merge into one analysis table (one row per ID)  ->  reports + staging CSVs + SPSS syntax.

Two phases so a person can review the plan:
    tables = load_tables(paths)              # read + type every table
    plan   = make_plan(tables)               # base table, link columns, grain rules (editable)
    result = execute_plan(tables, plan, out_dir)
"""
import csv, datetime as dt, difflib, hashlib, io, os, re
from collections import Counter, OrderedDict
from dataclasses import dataclass, field
import numpy as np
import pandas as pd
import openpyxl

MAX_EMPTY_STREAK = 5000
DATE_FORMATS = ["%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%d", "%d-%m-%Y", "%m-%d-%Y", "%d.%m.%Y", "%Y/%m/%d",
                "%d-%b-%Y", "%d %b %Y", "%d-%b-%y", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M", "%m/%d/%Y %H:%M"]
ID_NAME = re.compile(r"(^|[_\s])(id|key|code|no|num|number|uid|uhid|pid|hhid)$|record|participant|subject|patient|"
                     r"household|serial|sl_?no|registration|enrol|barcode|mrn", re.I)
NOT_ID_NAME = re.compile(r"name|date|age|score|time|amount|income|weight|height|comment|remark|note", re.I)
SEQ_NAME = re.compile(r"visit|round|wave|event|instance|follow|timepoint|time_point|month|week|phase|trimester|repeat", re.I)
BASE_HINT = re.compile(r"master|participant|registration|enrol|roster|demograph|baseline_id|id_list|screening", re.I)
PII_NAME = re.compile(r"(^|_)(name|husband|father|mother|spouse|phone|mobile|contact|address|email|aadhaar|aadhar|"
                      r"pan|voter|guardian)(_|$)|_name$|^name", re.I)


class PlanError(Exception):
    pass


# ---------------------------------------------------------------------------------------------- reading
def _md5(path):
    return hashlib.md5(open(path, "rb").read()).hexdigest()


def _read_csv_rows(path):
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc); break
        except UnicodeDecodeError:
            continue
    sample = text[:50000]
    try:
        delim = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        delim = max([",", ";", "\t", "|"], key=sample.count)
    rows = [[(c if c.strip() != "" else None) for c in r] for r in csv.reader(io.StringIO(text), delimiter=delim)]
    return [("", rows, len(rows))], delim, enc


def _read_xlsx_sheets(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows, streak = [], 0
        for r in ws.iter_rows(values_only=True):
            r = [None if (isinstance(v, str) and v.strip() == "") else v for v in r]
            if all(v is None for v in r):
                streak += 1
                if streak > MAX_EMPTY_STREAK and rows:
                    break
                rows.append(r); continue
            streak = 0; rows.append(r)
        while rows and all(v is None for v in rows[-1]):
            rows.pop()
        out.append((ws.title, rows, ws.max_row or len(rows)))
    wb.close()
    return out


def _col_letter(i):
    return openpyxl.utils.get_column_letter(i + 1)


def _find_header(rows):
    """First row (within 15) that looks like a header: wide enough and mostly non-numeric text."""
    width = max((sum(v is not None for v in r) for r in rows[:50]), default=0)
    for i, r in enumerate(rows[:15]):
        vals = [v for v in r if v is not None]
        if len(vals) < max(1, 0.5 * width):
            continue
        texty = sum(isinstance(v, str) and not _is_num_text(v) for v in vals)
        if texty >= 0.7 * len(vals):
            return i
    return 0


def _is_num_text(s):
    return isinstance(s, str) and bool(re.fullmatch(r"\s*[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?\s*", s))


def clean_text(v):
    s = str(v).replace("\t", " ").replace("\r", " ").replace("\n", " ").strip()
    return s or None


# ---------------------------------------------------------------------------------------------- typing
def _parse_date_text(values):
    """Return (format, ambiguous) if one format parses >=98% of the texts."""
    uniq = list(dict.fromkeys(v.strip() for v in values))[:3000]
    best, best_rate = None, 0.0
    rates = {}
    for f in DATE_FORMATS:
        ok = 0
        for s in uniq:
            try:
                dt.datetime.strptime(s, f); ok += 1
            except ValueError:
                pass
        rates[f] = ok / len(uniq)
        if rates[f] > best_rate:
            best, best_rate = f, rates[f]
    if best_rate < 0.98:
        return None, False
    pair = {"%d/%m/%Y": "%m/%d/%Y", "%m/%d/%Y": "%d/%m/%Y", "%d-%m-%Y": "%m-%d-%Y", "%m-%d-%Y": "%d-%m-%Y"}
    ambiguous = best in pair and rates.get(pair[best], 0) >= 0.98
    if ambiguous:  # both orders fit: day-first is the default (India / UK convention)
        best = "%d/%m/%Y" if "/" in best else "%d-%m-%Y"
    return best, ambiguous


def infer_column(values):
    """values: raw cell values. Returns (kind, converted list, info dict)."""
    vals = [clean_text(v) if isinstance(v, str) else v for v in values]
    nn = [v for v in vals if v is not None]
    info = {}
    if not nn:
        return "empty", [None] * len(vals), info
    if all(isinstance(v, dt.time) for v in nn):
        return "time", vals, info
    if all(isinstance(v, (dt.datetime, dt.date)) for v in nn):
        has_time = any(isinstance(v, dt.datetime) and (v.hour or v.minute or v.second) for v in nn)
        if has_time:
            return "datetime", [v if v is not None else None for v in vals], info
        return "date", [None if v is None else (v.date() if isinstance(v, dt.datetime) else v) for v in vals], info
    isnum = lambda v: (isinstance(v, (int, float)) and not isinstance(v, bool)) or _is_num_text(v)
    if all(isnum(v) for v in nn):
        long_digits = any(isinstance(v, str) and len(re.sub(r"\D", "", v)) > 15 for v in nn)
        if not long_digits:
            info["text_numbers"] = sum(isinstance(v, str) for v in nn)
            return "num", [None if v is None else float(v) for v in vals], info
    texts = [v for v in nn if isinstance(v, str)]
    if texts and len(texts) == len([v for v in nn if not isinstance(v, (dt.date, dt.datetime))]):
        fmt, amb = _parse_date_text(texts)
        if fmt:
            out = []
            for v in vals:
                if v is None: out.append(None)
                elif isinstance(v, dt.datetime): out.append(v.date())
                elif isinstance(v, dt.date): out.append(v)
                else:
                    try: out.append(dt.datetime.strptime(v.strip(), fmt).date())
                    except ValueError: out.append(None)
            info["date_format"] = fmt; info["date_ambiguous"] = amb
            return "date", out, info
    out = [None if v is None else clean_text(v.strftime("%H:%M") if isinstance(v, dt.time) else
                                             v.strftime("%Y-%m-%d") if isinstance(v, (dt.date, dt.datetime)) else v) for v in vals]
    return "text", out, info


def norm_key(v):
    """'001001', 1001, 1001.0 -> '1001';  ' p-01 ' -> 'P-01'."""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, np.integer)):
        return str(int(v))
    if isinstance(v, (float, np.floating)):
        if np.isnan(v) or not float(v).is_integer():
            return None
        return str(int(v))
    if isinstance(v, (dt.date, dt.datetime, dt.time)):
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    if not s:
        return None
    if s.isdigit():
        return str(int(s))
    if re.fullmatch(r"\d+\.0+", s):
        return str(int(float(s)))
    return s.upper()


# ---------------------------------------------------------------------------------------------- tables
@dataclass
class Table:
    tid: str
    file: str
    sheet: str
    md5: str
    rows_in_sheet: int
    header_row: int
    empty_rows: int
    data: pd.DataFrame                    # typed values, original column names (cleaned)
    raw_cols: dict                        # column -> raw values (for key normalisation)
    kinds: dict
    info: dict
    original_names: dict                  # column -> header as written in the file
    candidates: list = field(default_factory=list)
    notes: list = field(default_factory=list)


def _slug(s):
    s = re.sub(r"[^0-9a-zA-Z]+", "_", s).strip("_").lower()
    return s or "table"


def load_tables(paths, progress=None):
    """Read and type every table in every file. Returns list[Table] and list of skipped files."""
    say = progress or (lambda i, n, t: None)
    tables, skipped, used_ids = [], [], set()
    for k, path in enumerate(paths):
        name = os.path.basename(path); ext = os.path.splitext(name)[1].lower()
        say(k, len(paths), name)
        try:
            if ext in (".csv", ".tsv", ".txt"):
                sheets, delim, enc = _read_csv_rows(path)
            elif ext in (".xlsx", ".xlsm"):
                sheets = _read_xlsx_sheets(path)
            else:
                skipped.append(dict(file=name, reason="unsupported file type (use .csv, .tsv, .txt, .xlsx, .xlsm)")); continue
        except Exception as e:
            skipped.append(dict(file=name, reason=f"could not be read: {type(e).__name__}: {e}")); continue
        h = _md5(path)
        multi = len([s for s in sheets if any(any(v is not None for v in r) for r in s[1])]) > 1
        for sheet, rows, dim in sheets:
            if not any(any(v is not None for v in r) for r in rows):
                continue
            hi = _find_header(rows)
            header = list(rows[hi]); body = rows[hi + 1:]
            width = max([len(header)] + [len(r) for r in body[:2000]])
            header += [None] * (width - len(header))
            body = [list(r) + [None] * (width - len(r)) for r in body]
            nonempty = [r for r in body if any(v is not None for v in r)]
            # drop columns with neither header nor data
            keep_idx = [i for i in range(width) if header[i] is not None or any(r[i] is not None for r in nonempty)]
            base = _slug(os.path.splitext(name)[0]) + (f"_{_slug(sheet)}" if multi else "")
            names, seen, orig = [], Counter(), {}
            for i in keep_idx:
                hdr = header[i]
                nm = clean_text(hdr) if hdr is not None else None
                nm = nm or (f"{base}_col_{_col_letter(i)}" if ext.startswith(".xls") else f"{base}_col_{i + 1}")
                seen[nm.lower()] += 1
                if seen[nm.lower()] > 1:
                    nm = f"{nm}_{seen[nm.lower()]}"
                names.append(nm); orig[nm] = "" if hdr is None else str(hdr)
            raw_cols = {nm: [r[i] for r in nonempty] for nm, i in zip(names, keep_idx)}
            data, kinds, info = {}, {}, {}
            for nm in names:
                kind, vals, inf = infer_column(raw_cols[nm])
                data[nm] = vals; kinds[nm] = kind; info[nm] = inf
            tid, j = base, 2
            while tid in used_ids:
                tid = f"{base}_{j}"; j += 1
            used_ids.add(tid)
            t = Table(tid=tid, file=name, sheet=sheet if multi else "", md5=f"{h}:{sheet}", rows_in_sheet=dim, header_row=hi + 1,
                      empty_rows=max(0, dim - hi - 1 - len(nonempty)), data=pd.DataFrame(data, columns=names),
                      raw_cols=raw_cols, kinds=kinds, info=info, original_names=orig)
            t.candidates = key_candidates(t)
            tables.append(t)
    say(len(paths), len(paths), "done")
    return tables, skipped


def key_candidates(t):
    n = len(t.data)
    out = []
    if n == 0:
        return out
    for pos, c in enumerate(t.data.columns):
        kind = t.kinds[c]
        if kind in ("date", "datetime", "time", "empty"):
            continue
        if kind == "num" and not all(float(v).is_integer() for v in t.data[c].dropna()):
            continue
        keys = [norm_key(v) for v in t.raw_cols[c]]
        nn = [k for k in keys if k is not None]
        if len(nn) < 0.5 * n:
            continue
        if kind == "text" and np.mean([len(k) for k in nn]) > 30:
            continue
        uniq = len(set(nn)) / len(nn)
        if uniq < 0.05:
            continue
        name_bonus = (1.0 if ID_NAME.search(c) else 0.0) - (0.8 if NOT_ID_NAME.search(c) and not ID_NAME.search(c) else 0.0)
        score = 3 * uniq + len(nn) / n + name_bonus + (0.3 if pos == 0 else 0)
        out.append(dict(column=c, unique_ratio=round(uniq, 4), filled_ratio=round(len(nn) / n, 4), score=round(score, 3),
                        keys=keys, keyset=set(nn)))
    return sorted(out, key=lambda d: -d["score"])[:8]


# ---------------------------------------------------------------------------------------------- planning
def _containment(a, b):
    return len(a & b) / len(a) if a else 0.0


def table_label(t):
    return t.file + (f" [{t.sheet}]" if t.sheet else "")


def make_plan(tables, base_tid=None, base_key=None):
    """Choose base table + key, link column per table, grain rule. Returns an editable dict."""

    return _make_plan_impl(tables, base_tid, base_key)


def base_options(tables):
    """Tables (and their columns) that could serve as the base: unique, complete ID column."""
    out = {}
    for t in tables:
        cols = [c["column"] for c in t.candidates if c["unique_ratio"] >= 0.999 and c["filled_ratio"] >= 0.95]
        if cols:
            out[t.tid] = cols
    return out


def _make_plan_impl(tables, base_tid=None, base_key=None):
    if not tables:
        raise PlanError("No readable tables were found in the uploaded files.")
    # --- base table: a unique, complete key whose values cover the other tables
    scored = []
    for t in tables:
        for c in t.candidates:
            if c["unique_ratio"] < 0.999 or c["filled_ratio"] < 0.95:
                continue
            cover = 0
            for o in tables:
                if o is t:
                    continue
                if any(_containment(oc["keyset"], c["keyset"]) >= 0.8 for oc in o.candidates):
                    cover += 1
            bonus = (2 if BASE_HINT.search(t.file) else 0) - (3 if "copy" in t.file.lower() else 0)
            scored.append((cover, bonus, len(c["keyset"]), c["score"], t.tid, c["column"]))
    if base_tid:
        cand = [s for s in scored if s[4] == base_tid]
        if not cand:
            raise PlanError(f"{base_tid} has no column with unique values, so it cannot be the base table.")
        scored = cand
    if base_key:
        cand = [x for x in scored if x[5] == base_key]
        if not cand:
            raise PlanError(f"'{base_key}' does not have unique, complete values, so it cannot be the ID column.")
        scored = cand
    if not scored:
        raise PlanError("No table has a column with unique, complete values to use as the participant ID. "
                        "Upload a list with one row per participant.")
    scored.sort(reverse=True)
    _, _, _, _, btid, bcol = scored[0]
    base = next(t for t in tables if t.tid == btid)
    bkeys = next(c for c in base.candidates if c["column"] == bcol)["keyset"]
    # --- order: originals before copies, then file name
    order = sorted([t for t in tables if t is not base], key=lambda t: ("copy" in t.file.lower(), t.file.lower(), t.sheet))
    items = [dict(tid=base.tid, label=table_label(base), role="base", link=bcol, link_options=[c["column"] for c in base.candidates],
                  match_pct=100.0, grain="one row per ID", rule="direct", rule_options=["direct"], include=True,
                  rows=len(base.data), ids=len(bkeys), note="")]
    for t in order:
        best = None
        for c in t.candidates:
            cont = _containment(c["keyset"], bkeys)
            rank = (round(cont, 3), c["column"].lower() == bcol.lower(), c["unique_ratio"], c["score"])
            if best is None or rank > best[0]:
                best = (rank, c)
        link_options = [c["column"] for c in t.candidates]
        if best is None or best[0][0] < 0.5:
            items.append(dict(tid=t.tid, label=table_label(t), role="unlinked", link=None, link_options=link_options,
                              match_pct=round(100 * best[0][0], 1) if best else 0.0, grain="-", rule="exclude",
                              rule_options=["exclude"], include=False, rows=len(t.data), ids=0,
                              note="no column matches the base IDs (needs at least 50% of values to match)"))
            continue
        c = best[1]
        gi = _grain_item(t, c, round(100 * best[0][0], 1))
        same = [x for x in tables if x is not t and x.md5 == t.md5 and (x is base or order.index(x) < order.index(t))]
        if same:
            gi["note"] = f"identical copy of {table_label(same[0])} - will be skipped"
        items.append(gi)
    return dict(base=base.tid, key=bcol, items=items)


def _grain_item(t, c, match_pct):
    keys = pd.Series(c["keys"])
    nn = keys.dropna()
    unique = nn.is_unique
    options, rule, note = ["direct"], "direct", ""
    if not unique:
        dup_ids = int(nn.duplicated(keep=False).groupby(nn).any().sum())
        options = []
        seq = []
        for col in t.data.columns:
            if col == c["column"] or t.kinds[col] in ("empty",):
                continue
            lv = t.data[col].dropna().unique()
            if 2 <= len(lv) <= 12:
                combo = pd.DataFrame({"k": keys, "v": t.data[col]}).dropna()
                if not combo.duplicated().any() and len(combo) >= 0.95 * len(nn):
                    seq.append((bool(SEQ_NAME.search(col)), -len(lv), col))
        seq.sort(reverse=True)
        options += [f"one column per {s[2]}" for s in seq[:3]]
        dates = [col for col in t.data.columns if t.kinds[col] in ("date", "datetime")]
        options += [f"latest by {d}" for d in dates[:3]]
        options += ["first row", "mean of numbers", "keep as separate table"]
        rule = options[0]
        note = f"{dup_ids} IDs have more than one row"
    return dict(tid=t.tid, label=table_label(t), role="linked", link=c["column"],
                link_options=[x["column"] for x in t.candidates], match_pct=match_pct,
                grain="one row per ID" if unique else "several rows per ID", rule=rule, rule_options=options,
                include=True, rows=len(t.data), ids=int(nn.nunique()), note=note)


def update_item(tables, plan, tid, new_link):
    """Re-evaluate one table after the user picks a different link column."""
    t = next(x for x in tables if x.tid == tid)
    base = next(x for x in tables if x.tid == plan["base"])
    bkeys = next(c for c in base.candidates if c["column"] == plan["key"])["keyset"]
    c = next(x for x in t.candidates if x["column"] == new_link)
    item = _grain_item(t, c, round(100 * _containment(c["keyset"], bkeys), 1))
    for i, it in enumerate(plan["items"]):
        if it["tid"] == tid:
            plan["items"][i] = item
    return item


# ---------------------------------------------------------------------------------------------- execution
def _letters(s):
    return re.sub(r"[^a-z]", "", str(s).lower()) if s is not None else ""


def _same(a, b):
    if a is None or b is None:
        return True
    if isinstance(a, float) and isinstance(b, float):
        return abs(a - b) < 1e-9
    if isinstance(a, str) and isinstance(b, str):
        return a.strip().lower() == b.strip().lower()
    return a == b


def spss_name(name, used):
    s = re.sub(r"[^0-9A-Za-z_.@#$]", "_", str(name))
    s = re.sub(r"_+", "_", s).strip("_.") or "var"
    if not re.match(r"[A-Za-z@]", s):
        s = "v" + s
    if s.upper() in {"ALL", "AND", "BY", "EQ", "GE", "GT", "LE", "LT", "NE", "NOT", "OR", "TO", "WITH"}:
        s = s + "_"
    s = s[:64]
    base, i = s, 2
    while s.lower() in used:
        suf = f"_{i}"; s = base[:64 - len(suf)] + suf; i += 1
    used.add(s.lower())
    return s


def _apply_rule(df, rule, kinds):
    """df has '__key'. Returns (one-row-per-key frame, new kinds, note)."""
    if rule == "direct":
        if df["__key"].duplicated().any():
            df = df.drop_duplicates("__key", keep="first")
            return df, kinds, "duplicate IDs found - first row kept"
        return df, kinds, ""
    if rule.startswith("one column per "):
        col = rule[len("one column per "):]
        levels = sorted(df[col].dropna().unique(), key=lambda v: (str(type(v)), v))
        parts, nk = [], {}
        for lv in levels:
            sub = df[df[col] == lv].drop(columns=[col]).drop_duplicates("__key")
            suffix = _slug(str(int(lv)) if isinstance(lv, float) and float(lv).is_integer() else str(lv))
            ren = {c: f"{c}_{suffix}" for c in sub.columns if c != "__key"}
            nk.update({ren[c]: kinds[c] for c in ren})
            parts.append(sub.rename(columns=ren).set_index("__key"))
        out = pd.concat(parts, axis=1).reset_index()
        return out, nk, f"{len(levels)} levels of {col} spread into columns"
    if rule.startswith("latest by "):
        col = rule[len("latest by "):]
        o = df.assign(__ord=pd.to_datetime(df[col], errors="coerce")).sort_values("__ord", na_position="first", kind="stable")
        return o.drop_duplicates("__key", keep="last").drop(columns="__ord"), kinds, f"latest row by {col} kept"
    if rule == "first row":
        return df.drop_duplicates("__key", keep="first"), kinds, "first row per ID kept"
    if rule == "mean of numbers":
        g = df.groupby("__key", sort=False)
        cols = {}
        for c in df.columns:
            if c == "__key": continue
            if kinds[c] == "num": cols[c] = g[c].mean()
            elif kinds[c] in ("date", "datetime"): cols[c] = g[c].max()
            else: cols[c] = g[c].first()
        out = pd.DataFrame(cols); out["n_rows"] = g.size(); out = out.reset_index()
        nk = dict(kinds); nk["n_rows"] = "num"
        return out, nk, "numbers averaged, dates = latest, text = first value; n_rows added"
    raise PlanError(f"Unknown rule: {rule}")


def execute_plan(tables, plan, out_dir, drop_empty=True, progress=None):
    say = progress or (lambda i, n, t: None)
    os.makedirs(out_dir, exist_ok=True)
    T = {t.tid: t for t in tables}
    base = T[plan["base"]]; bcol = plan["key"]
    items = {it["tid"]: it for it in plan["items"]}
    log, profile, qc, not_merged, separate = [], [], [], [], {}
    kinds, labels, source, owner = {}, {}, {}, {}

    # ---- base
    bkeys_raw = base.raw_cols[bcol]
    bkeys = [norm_key(v) for v in bkeys_raw]
    if any(k is None for k in bkeys) or len(set(bkeys)) != len(bkeys):
        raise PlanError(f"The base table's ID column '{bcol}' has missing or repeated values.")
    numeric_key = all(k.isdigit() for k in bkeys)
    merged = base.data.copy()
    merged.insert(0, "__key", bkeys)
    merged[bcol] = [int(k) for k in bkeys] if numeric_key else bkeys
    kinds[bcol] = "num" if numeric_key else "text"
    for c in base.data.columns:
        if c == bcol: continue
        if drop_empty and base.kinds[c] == "empty":
            merged = merged.drop(columns=c); log.append(dict(table=base.tid, column=c, decision="dropped", detail="empty column"))
            continue
        kinds[c] = base.kinds[c]
    for c in merged.columns:
        if c == "__key": continue
        source[c] = (table_label(base), base.original_names.get(c, c)); owner[c] = base.tid
        labels[c] = base.original_names.get(c, c) or c
    base_ids = set(bkeys)
    base_text = [c for c in merged.columns if c not in ("__key", bcol) and kinds.get(c) == "text"]
    presence = {base.tid: pd.Series(True, index=range(len(merged)))}
    accepted = {base.tid: dict(keys=base_ids, cols={c: merged.set_index("__key")[c] for c in merged.columns if c != "__key"})}
    seen_md5 = {base.md5: table_label(base)}
    profile.append(_profile_row(base, items[base.tid], bcol, len(base.data), 0, 0, 0, "base table"))
    leading_zero_widths = Counter(len(v.strip()) for v in bkeys_raw if isinstance(v, str) and v.strip().isdigit() and v.strip().startswith("0"))

    order = [it for it in plan["items"] if it["tid"] != base.tid]
    for n, it in enumerate(order, 1):
        t = T[it["tid"]]
        say(n, len(order), table_label(t))
        if not it.get("include", True) or it["rule"] == "exclude" or not it.get("link"):
            not_merged.append(dict(file=table_label(t), reason=it.get("note") or "excluded by user"))
            profile.append(_profile_row(t, it, it.get("link"), 0, 0, 0, 0, "not merged: " + (it.get("note") or "excluded")))
            continue
        if t.md5 in seen_md5:
            not_merged.append(dict(file=table_label(t), reason=f"byte-identical copy of {seen_md5[t.md5]}"))
            profile.append(_profile_row(t, it, it["link"], 0, 0, 0, 0, f"duplicate of {seen_md5[t.md5]}"))
            continue
        link = it["link"]
        raw = t.raw_cols[link]
        leading_zero_widths.update(len(v.strip()) for v in raw if isinstance(v, str) and v.strip().isdigit() and v.strip().startswith("0"))
        df = t.data.drop(columns=[link]).copy()
        df.insert(0, "__key", [norm_key(v) for v in raw])
        tk = {c: t.kinds[c] for c in df.columns if c != "__key"}
        no_key = int(df["__key"].isna().sum()); df = df[df["__key"].notna()]
        orphan = int((~df["__key"].isin(base_ids)).sum()); df = df[df["__key"].isin(base_ids)]
        exact = int(df.astype(str).duplicated().sum()); df = df[~df.astype(str).duplicated()]
        if it["rule"] == "keep as separate table":
            sep = df.copy(); sep.insert(0, bcol, [int(k) if numeric_key else k for k in sep["__key"]])
            separate[t.tid] = (sep.drop(columns="__key"), {**{bcol: kinds[bcol]}, **tk}, table_label(t))
            not_merged.append(dict(file=table_label(t), reason="kept as a separate long table (several rows per ID)"))
            profile.append(_profile_row(t, it, link, len(df), no_key, orphan, exact, "separate table"))
            continue
        df, tk, rule_note = _apply_rule(df.reset_index(drop=True), it["rule"], tk)
        tkeys = set(df["__key"])
        # ---- content duplicate of an accepted table?
        dup_of = None
        for atid, a in accepted.items():
            if a["keys"] != tkeys:
                continue
            common = [c for c in df.columns if c != "__key" and c in a["cols"]]
            if not common:
                continue
            s = df.set_index("__key")
            if all(all(_same(x, y) for x, y in zip(s[c].reindex(a["cols"][c].index), a["cols"][c])) for c in common):
                rest = [c for c in df.columns if c not in common and c != "__key"]
                if all(tk.get(c) == "empty" or _copy_of_base(s[c], merged, base_text)[0] == "copy" for c in rest):
                    dup_of = atid; break
        if dup_of:
            not_merged.append(dict(file=table_label(t), reason=f"same IDs and values as {table_label(T[dup_of])}"))
            profile.append(_profile_row(t, it, link, len(df), no_key, orphan, exact, f"same data as {table_label(T[dup_of])}"))
            continue
        # ---- column decisions
        s = df.set_index("__key")
        add, counts = [], Counter()
        mk = merged.set_index("__key")
        for c in [c for c in df.columns if c != "__key"]:
            col = s[c]
            if tk.get(c) == "empty" or col.isna().all():
                if drop_empty:
                    counts["empty"] += 1; log.append(dict(table=t.tid, column=c, decision="dropped", detail="empty column")); continue
            if tk.get(c) == "text":
                strength, bc, mism = _copy_of_base(col, merged, base_text)
                for k, a, b, sim in mism:
                    qc.append(dict(id=k, table=table_label(t), column=c, base_column=bc, base_value=a, file_value=b, similarity=sim))
                if strength == "copy":
                    counts["copy"] += 1
                    log.append(dict(table=t.tid, column=c, decision="dropped", detail=f"copy of base column {bc} (>=95% agree); {len(mism)} mismatches listed in consistency check"))
                    continue
                if strength == "related":
                    log.append(dict(table=t.tid, column=c, decision="kept (related)", detail=f"similar to base column {bc} but <95% agree; kept, {len(mism)} mismatches listed in consistency check"))
                    if c in mk.columns:
                        new = f"{c}__{t.tid}"; s = s.rename(columns={c: new}); tk[new] = tk.pop(c)
                        t.original_names[new] = t.original_names.get(c, c); c = new; counts["conflict"] += 1
                        add.append(c); counts["added"] += 1
                        source[c] = (table_label(t), t.original_names.get(c, c.split("__")[0]))
                        labels[c] = f"{t.original_names.get(c, c) or c} [{t.file}]"; kinds[c] = tk[c]; owner[c] = t.tid
                        continue
            if c in mk.columns:
                ov = mk[c].reindex(col.index)
                both = ov.notna() & col.notna()
                agree = sum(_same(x, y) for x, y in zip(ov[both], col[both]))
                if both.sum() > 0 and agree / both.sum() >= 0.995:
                    fill = col[ov.isna() & col.notna()]
                    if len(fill):
                        merged.loc[merged["__key"].isin(fill.index), c] = merged.loc[merged["__key"].isin(fill.index), "__key"].map(fill)
                    counts["shared"] += 1
                    log.append(dict(table=t.tid, column=c, decision="kept once", detail=f"same values as existing column ({agree}/{int(both.sum())} agree); {len(fill)} missing values filled"))
                    continue
                new = f"{c}__{t.tid}"
                counts["conflict"] += 1
                log.append(dict(table=t.tid, column=c, decision="renamed", detail=f"name already used with different values ({agree}/{int(both.sum())} agree) -> {new}"))
                s = s.rename(columns={c: new}); tk[new] = tk.pop(c); t.original_names[new] = t.original_names.get(c, c); c = new
            add.append(c); counts["added"] += 1
            source[c] = (table_label(t), t.original_names.get(c, c.split("__")[0]))
            labels[c] = f"{t.original_names.get(c, c) or c} [{t.file}]" if c.endswith("__" + t.tid) else (t.original_names.get(c, c) or c)
            kinds[c] = tk[c]; owner[c] = t.tid
            if c not in [x for x in tk]: pass
        if add:
            merged = merged.merge(s[add].reset_index(), on="__key", how="left", validate="one_to_one")
        presence[t.tid] = merged["__key"].isin(tkeys)
        accepted[t.tid] = dict(keys=tkeys, cols={c: s[c] for c in s.columns})
        seen_md5[t.md5] = table_label(t)
        st = f"merged ({it['rule']}{': ' + rule_note if rule_note else ''})"
        p = _profile_row(t, it, link, len(df), no_key, orphan, exact, st)
        p.update(columns_added=counts["added"], columns_kept_once=counts["shared"], columns_renamed=counts["conflict"],
                 copy_columns_dropped=counts["copy"], empty_columns_dropped=counts["empty"])
        profile.append(p)

    # ---- presence flags, padded id
    flag_cols = []
    for tid, pres in presence.items():
        nm = "has_" + tid[:56]; merged[nm] = pres.astype(int).values; kinds[nm] = "num"; flag_cols.append(nm)
        labels[nm] = f"Has data in {table_label(T[tid])}"; source[nm] = ("derived", ""); owner[nm] = "_flags"
    merged["n_datasets"] = merged[flag_cols].sum(axis=1); kinds["n_datasets"] = "num"; labels["n_datasets"] = "Number of datasets with data"
    source["n_datasets"] = ("derived", ""); owner["n_datasets"] = "_flags"
    if numeric_key and leading_zero_widths:
        w = leading_zero_widths.most_common(1)[0][0]
        nm = f"{bcol}_padded"
        merged.insert(2, nm, [str(k).zfill(w) for k in merged["__key"]]); kinds[nm] = "text"
        labels[nm] = f"{bcol} as {w}-digit text with leading zeros"; source[nm] = ("derived", ""); owner[nm] = base.tid
    # ---- SPSS-safe names
    used, rename = set(), {}
    for c in merged.columns:
        if c == "__key": continue
        rename[c] = spss_name(c, used)
    merged = merged.rename(columns=rename)
    R = lambda d: {rename.get(k, k): v for k, v in d.items()}
    kinds, labels, source, owner = R(kinds), R(labels), R(source), R(owner)
    log = [dict(l, column_in_output=rename.get(l["column"], "")) for l in log]
    key_out = rename[bcol]
    # ---- outliers
    outliers = []
    for c in merged.columns:
        if kinds.get(c) != "num" or c in (key_out, "n_datasets") or c.startswith("has_"):
            continue
        x = merged[c].astype(float); nn = x.dropna()
        if len(nn) < 30 or nn.nunique() < 10:
            continue
        med = nn.median(); mad = (nn - med).abs().median()
        if mad == 0:
            continue
        z = 0.6745 * (x - med) / mad
        for i in z[z.abs() > 7].index[:200]:
            outliers.append(dict(id=merged.at[i, key_out], column=c, value=x[i], median=med, robust_z=round(float(z[i]), 1)))
    # ---- dictionary
    order_cols = [key_out] + [c for c in merged.columns if c not in ("__key", key_out)]
    final = merged[order_cols]
    dictionary = pd.DataFrame([dict(variable=c, label=labels.get(c, ""), type=kinds.get(c), source_file=source.get(c, ("", ""))[0],
                                    original_column=source.get(c, ("", ""))[1], n_filled=int(final[c].notna().sum())) for c in final.columns])
    result = dict(merged=final, key=key_out, kinds=kinds, labels=labels, source=source, owner=owner, dictionary=dictionary,
                  profile=pd.DataFrame(profile), log=pd.DataFrame(log), consistency=pd.DataFrame(qc), outliers=pd.DataFrame(outliers),
                  not_merged=pd.DataFrame(not_merged), separate=separate, out_dir=out_dir, numeric_key=numeric_key,
                  base_label=table_label(base))
    write_staging_and_syntax(result)
    say(len(order), len(order), "done")
    return result


def _copy_of_base(col, merged, base_text):
    """Is this text column a copy of a base text column (e.g. a repeated participant name)?"""
    col = col.dropna()
    if len(col) < 20 or not base_text:
        return False, None, []
    mk = merged.set_index("__key")
    best = (0, None)
    sample = col.index[:3000]
    for bc in base_text:
        ref = mk[bc].reindex(sample)
        ok = n = 0
        for a, b in zip(ref, col.reindex(sample)):
            if a is None or (isinstance(a, float) and a != a):
                continue
            n += 1
            la, lb = _letters(a), _letters(b)
            if la == lb or (la and lb and difflib.SequenceMatcher(None, la, lb).ratio() >= 0.85):
                ok += 1
        if n >= 20 and ok / n > best[0]:
            best = (ok / n, bc)
    if best[0] < 0.7:
        return False, None, []
    bc = best[1]; ref = mk[bc].reindex(col.index)
    strength = "copy" if best[0] >= 0.95 else "related"
    mism = []
    for k, a, b in zip(col.index, ref, col):
        if a is None or (isinstance(a, float) and a != a):
            continue
        la, lb = _letters(a), _letters(b)
        if la == lb:
            continue
        sim = difflib.SequenceMatcher(None, la, lb).ratio() if la and lb else 0.0
        if sim < 0.6:
            mism.append((k, a, b, round(sim, 2)))
    return strength, bc, mism


def _profile_row(t, it, link, rows_used, no_key, orphan, exact, status):
    raw = t.raw_cols.get(link, []) if link else []
    return OrderedDict(
        table=table_label(t), role=it.get("role"), rows_in_sheet=t.rows_in_sheet, header_row=t.header_row,
        data_rows=len(t.data), empty_rows_dropped=t.empty_rows, columns=len(t.data.columns), link_column=link,
        id_text_cells=sum(isinstance(v, str) for v in raw), id_number_cells=sum(isinstance(v, (int, float)) for v in raw),
        rows_without_id=no_key, ids_not_in_base=orphan, exact_duplicate_rows=exact, match_pct=it.get("match_pct"),
        grain=it.get("grain"), rule=it.get("rule"), rows_after_rule=rows_used,
        text_numbers_converted=sum(i.get("text_numbers", 0) for i in t.info.values()),
        date_formats="; ".join(sorted({f"{c}: {i['date_format']}{' (ambiguous - day first assumed)' if i.get('date_ambiguous') else ''}"
                                       for c, i in t.info.items() if i.get("date_format")})),
        status=status)


# ---------------------------------------------------------------------------------------------- staging + SPSS
def _fmt_num(x):
    if x is None or (isinstance(x, float) and np.isnan(x)): return ""
    x = float(x)
    return str(int(x)) if x.is_integer() else np.format_float_positional(x, trim="-")


def _cell(v, kind):
    if v is None or v is pd.NaT or (isinstance(v, float) and np.isnan(v)):
        return ""
    if kind == "date": return v.strftime("%m/%d/%Y")
    if kind == "datetime": return pd.Timestamp(v).strftime("%d-%b-%Y %H:%M:%S").upper()
    if kind == "time": return v.strftime("%H:%M:%S")
    if kind == "num": return _fmt_num(v)
    return str(v)


def write_staging_and_syntax(r):
    m, key, kinds, owner = r["merged"], r["key"], r["kinds"], r["owner"]
    stg = os.path.join(r["out_dir"], "staging"); os.makedirs(stg, exist_ok=True)
    groups = OrderedDict()
    for c in m.columns:
        if c == key: continue
        groups.setdefault(owner.get(c, "_other"), []).append(c)
    files = []
    for i, (g, cols) in enumerate(groups.items()):
        sub = m[[key] + cols]
        if i > 0:
            sub = sub[sub[cols].notna().any(axis=1)]
        fn = f"{i:02d}_{g[:40]}.csv"
        with open(os.path.join(stg, fn), "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh); w.writerow(sub.columns)
            for row in sub.itertuples(index=False):
                w.writerow([_cell(v, kinds.get(c, "text")) for c, v in zip(sub.columns, row)])
        files.append((fn, [key] + cols))
    for t, (df, k2, label) in r["separate"].items():
        df.to_csv(os.path.join(stg, f"separate_{t[:40]}.csv"), index=False)
    r["staging_files"] = files
    open(os.path.join(r["out_dir"], "merge_syntax.sps"), "w", encoding="utf-8").write(make_spss_syntax(r))


def make_spss_syntax(r, cd="C:\\analysis\\merge_package"):
    m, key, kinds = r["merged"], r["key"], r["kinds"]
    def width(c):
        s = m[c].dropna().astype(str)
        return max(8, int(s.map(lambda x: len(x.encode("utf-8"))).max()) + 10) if len(s) else 8
    keyfmt = "F20.0" if r["numeric_key"] else f"A{width(key)}"
    fmt_in = {"num": "F20.0", "date": "ADATE10", "datetime": "DATETIME20", "time": "TIME8"}
    fmt_out = {"date": "DATE11", "datetime": "DATETIME20", "time": "TIME5"}
    L = [f"""* Encoding: UTF-8.
* =====================================================================.
* Analysis file - SPSS rebuild from the staging CSVs.
* Generated by the Analysis File Builder. Base table: {r['base_label']}. Key: {key}.
* One staging CSV per source table; every column type is declared, so the result does not
* depend on SPSS guessing types. Change the CD line to the folder that contains 'staging'.
* The VERIFICATION table at the end must say PASS on every row.
* =====================================================================.
CD '{cd}'.
"""]
    tmp = []
    for i, (fn, cols) in enumerate(r["staging_files"]):
        spec = "\n".join(f"    {c} {keyfmt if c == key else fmt_in.get(kinds.get(c), None) or 'A' + str(width(c))}" for c in cols)
        L.append(f"* ---- {fn}.\nGET DATA /TYPE=TXT /FILE='staging/{fn}' /ENCODING='UTF8'\n  /DELCASE=LINE /DELIMITERS=\",\" /QUALIFIER='\"'\n"
                 f"  /ARRANGEMENT=DELIMITED /FIRSTCASE=2\n  /VARIABLES=\n{spec}.")
        disp = [f"{c} ({fmt_out[kinds[c]]})" for c in cols if kinds.get(c) in fmt_out]
        nums = [c for c in cols if kinds.get(c) == "num"]
        ints = [c for c in nums if m[c].dropna().map(lambda x: float(x).is_integer()).all()]
        decs = [c for c in nums if c not in ints]
        lines = []
        if ints: lines.append(" ".join(ints) + " (F8.0)")
        if decs: lines.append(" ".join(decs) + " (F10.2)")
        lines += disp
        if lines:
            L.append("FORMATS\n  " + "\n  ".join(lines) + ".")
        t = f"_tmp_{i:02d}.sav"; tmp.append(t)
        L.append(f"SORT CASES BY {key} (A).\nSAVE OUTFILE='{t}'.\n")
    L.append("MATCH FILES\n" + "\n".join(f"  /FILE='{t}'" + (" /IN=in_base" if j == 0 else "") for j, t in enumerate(tmp)) + f"\n  /BY {key}.")
    L.append("SELECT IF (in_base = 1).\nEXECUTE.\nDELETE VARIABLES in_base.")
    order = " ".join(m.columns)
    L.append(f"ADD FILES /FILE=* /KEEP={order}.\nEXECUTE.\nDATASET NAME merged.")
    labs = [(c, str(r["labels"].get(c) or "").replace("'", "")[:250]) for c in m.columns if r["labels"].get(c)]
    if labs:
        L.append("VARIABLE LABELS\n" + "\n".join(f"  {c} '{l}'" for c, l in labs) + ".")
    checks = [key] + [c for c in m.columns if c != key and kinds.get(c) in ("num", "date", "datetime", "time")]
    L.append("""* ---- VERIFICATION: filled values per variable vs the counts from the builder. All rows must say PASS.
DATASET ACTIVATE merged.
COMPUTE zz_one = 1.
DATASET DECLARE verify.
AGGREGATE /OUTFILE=verify /BREAK=zz_one""")
    L.append("\n".join(f"  /c{j:04d} = NU({c})" for j, c in enumerate(checks, 1)) + ".")
    L.append("DATASET ACTIVATE verify.\nSTRING zz_name (A8).\nCOMPUTE zz_name = 'found'.\nEXECUTE.\nDELETE VARIABLES zz_one.\nFLIP /NEWNAMES=zz_name.\n"
             "DELETE VARIABLES CASE_LBL.\nCOMPUTE check_no = $CASENUM.\nFORMATS check_no (F5.0).\nNUMERIC expected (F8.0).\nSTRING check_variable (A64).\nEXECUTE.")
    L.append("\n".join(f"IF (check_no = {j}) expected = {int(m[c].notna().sum())}.\nIF (check_no = {j}) check_variable = '{c}'." for j, c in enumerate(checks, 1)))
    L.append("""STRING result (A4).
COMPUTE result = 'FAIL'.
IF (found = expected) result = 'PASS'.
EXECUTE.
SORT CASES BY result (A) check_no (A).
LIST check_no check_variable expected found result.
FREQUENCIES VARIABLES=result.
DATASET ACTIVATE merged.
DELETE VARIABLES zz_one.
DATASET CLOSE verify.
SAVE OUTFILE='analysis_file.sav' /COMPRESSED.
SAVE TRANSLATE OUTFILE='analysis_file.xlsx' /TYPE=XLS /VERSION=12 /MAP /FIELDNAMES VALUE=NAMES /REPLACE.""")
    L.append("\n".join(f"ERASE FILE='{t}'." for t in tmp))
    return "\n".join(L) + "\n"
